import io
import sqlite3
import zipfile

import pytest
from docx import Document

from tests.conftest import SECRETS
from transcriber.export import export_answer, export_transcript
from transcriber.importers import import_bytes
from transcriber.mask import detect_and_assign, normalize_transcript
from transcriber.model import Transcript, Utterance
from transcriber.store import Store, WrongPassword
from transcriber.vault import Vault

# ---------------------------------------------------------------- импорт


def test_import_speaker_lines():
    data = "Оператор: Добрый день.\nКлиент: Здравствуйте,\nпродолжение реплики.\n[00:12:30] Иван: Да.\n"
    utts = import_bytes(data.encode(), "call.txt")
    assert [(u.speaker, u.text) for u in utts] == [
        ("Оператор", "Добрый день."),
        ("Клиент", "Здравствуйте, продолжение реплики."),
        ("Иван", "Да."),
    ]
    assert utts[2].start == 12 * 60 + 30


def test_import_cp1251_and_paragraphs():
    data = "Первый абзац без спикеров.\n\nВторой абзац, номер 123.\n".encode("cp1251")
    utts = import_bytes(data, "notes.txt")
    assert [u.text for u in utts] == ["Первый абзац без спикеров.", "Второй абзац, номер 123."]
    assert all(u.speaker is None for u in utts)


def test_import_srt():
    srt = "1\n00:00:01,000 --> 00:00:03,500\nСпикер 1: Привет\n\n2\n00:00:04,000 --> 00:00:05,000\nСпикер 1: ещё\n"
    utts = import_bytes(srt.encode(), "a.srt")
    assert len(utts) == 1
    assert utts[0].speaker == "Спикер 1"
    assert utts[0].text == "Привет ещё"
    assert utts[0].start == 1.0


def test_import_vtt_voice():
    vtt = "WEBVTT\n\n00:01.000 --> 00:02.000\n<v Анна>Добрый день</v>\n\n00:03.000 --> 00:04.000\n<v Клиент>Да\n"
    utts = import_bytes(vtt.encode(), "a.vtt")
    assert [(u.speaker, u.text) for u in utts] == [("Анна", "Добрый день"), ("Клиент", "Да")]


def test_import_docx():
    doc = Document()
    doc.add_paragraph("Оператор: Как вас зовут?")
    doc.add_paragraph("Клиент: Петров Иван.")
    buf = io.BytesIO()
    doc.save(buf)
    utts = import_bytes(buf.getvalue(), "t.docx")
    assert [u.speaker for u in utts] == ["Оператор", "Клиент"]


def test_import_unsupported():
    with pytest.raises(ValueError):
        import_bytes(b"x", "a.pdf")


# ---------------------------------------------------------------- экспорт


@pytest.fixture
def reviewed(call_transcript):
    vault = Vault(code="DEL-AB12")
    normalize_transcript(call_transcript)
    detect_and_assign(call_transcript, vault)
    call_transcript.reviewed = True
    return call_transcript, vault


def test_export_requires_review(call_transcript, tmp_path):
    with pytest.raises(PermissionError):
        export_transcript(call_transcript, Vault(code="DEL-0000"), tmp_path)


def test_docx_has_no_pii_anywhere(reviewed, tmp_path):
    tr, vault = reviewed
    path = export_transcript(tr, vault, tmp_path, "docx")
    assert path.name == "DEL-AB12_транскрипт.docx"
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("word/theme"):
                continue
            data = z.read(name).decode("utf-8", "ignore").lower()
            for s in SECRETS:
                assert s.lower() not in data, (s, name)
        core = z.read("docProps/core.xml").decode()
    assert "python-docx" not in core
    assert "<dc:creator></dc:creator>" in core or "<dc:creator/>" in core
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Не раскрывай и не переписывай метки" in text
    assert "DEL-AB12" in text


def test_md_export_and_unique_names(reviewed, tmp_path):
    tr, vault = reviewed
    a = export_transcript(tr, vault, tmp_path, "md")
    b = export_transcript(tr, vault, tmp_path, "md")
    assert a.name != b.name
    assert "[ФИО_" in a.read_text(encoding="utf-8")
    assert not list(tmp_path.glob(".tmp_*"))


def test_answer_docx(tmp_path):
    path = export_answer("# Итог\n- **Петров** перезвонит\n1. шаг", "DEL-AB12", tmp_path)
    paras = [p.text for p in Document(str(path)).paragraphs]
    assert paras == ["Итог", "Петров перезвонит", "шаг"]


# ---------------------------------------------------------------- хранилище


def test_store_encrypts_everything(tmp_path):
    db = tmp_path / "cases.db"
    st = Store(db)
    st.init("correct horse battery")
    case = st.create_case("Секретное дело Петрова")
    tr = Transcript(id="tr1", title="Звонок Петрову", utterances=[Utterance("Петров Иван, телефон 89161234567")])
    detect_and_assign(tr, case.vault)
    st.save_case(case)
    st.save_transcript(case.id, tr)
    st.close()
    raw = db.read_bytes()
    for s in ("Петров".encode(), "Секретное".encode(), b"89161234567", "Звонок".encode()):
        assert s not in raw

    st2 = Store(db)
    with pytest.raises(WrongPassword):
        st2.unlock("wrong password!")
    st2.unlock("correct horse battery")
    loaded = st2.get_case(case.id)
    assert loaded.name == "Секретное дело Петрова"
    assert st2.find_case_by_code(case.code).id == case.id
    _, tr2 = st2.get_transcript("tr1")
    assert tr2.spans and tr2.spans[0].entity_id
    st2.delete_case(case.id)
    assert st2.list_cases() == []
    con = sqlite3.connect(db)
    assert con.execute("SELECT count(*) FROM transcripts").fetchone()[0] == 0


def test_store_short_password(tmp_path):
    with pytest.raises(ValueError):
        Store(tmp_path / "x.db").init("short")


def test_purge_expired(tmp_path):
    st = Store(tmp_path / "x.db")
    st.init("password123")
    c = st.create_case("старое")
    st._db.execute("UPDATE cases SET updated=0 WHERE id=?", (c.id,))
    st.create_case("новое")
    assert st.purge_expired(30) == 1
    assert [x.name for x in st.list_cases()] == ["новое"]
