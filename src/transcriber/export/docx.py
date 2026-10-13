"""docx без ПД в метаданных."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

_EPOCH = dt.datetime(2000, 1, 1, tzinfo=dt.UTC)


def _clean_metadata(doc) -> None:
    cp = doc.core_properties
    for attr in ("author", "category", "comments", "content_status", "identifier", "keywords", "language",
                 "last_modified_by", "subject", "title", "version"):
        setattr(cp, attr, "")
    cp.revision = 1
    cp.created = _EPOCH
    cp.modified = _EPOCH
    cp.last_printed = _EPOCH


def _new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    _clean_metadata(doc)
    return doc


def transcript_docx(code: str, lines: list[str], path: Path) -> None:
    from transcriber.export import INSTRUCTION, PHONE_EMAIL_NOTE

    doc = _new_doc()
    doc.add_heading(f"Транскрипт {code}", level=1)
    p = doc.add_paragraph()
    p.add_run(INSTRUCTION).italic = True
    p = doc.add_paragraph()
    p.add_run(PHONE_EMAIL_NOTE).italic = True
    for line in lines:
        head, sep, body = line.partition(": ")
        para = doc.add_paragraph()
        if sep and len(head) <= 60:
            para.add_run(head + ":").bold = True
            para.add_run(" " + body)
        else:
            para.add_run(line)
    doc.save(str(path))


_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")


def _add_inline(para, text: str) -> None:
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith(("**", "__")) and len(chunk) > 4:
            para.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`"):
            para.add_run(chunk[1:-1]).font.name = "Courier New"
        elif chunk[0] in "*_" and chunk[-1] == chunk[0] and len(chunk) > 2:
            para.add_run(chunk[1:-1]).italic = True
        else:
            para.add_run(chunk)


def markdown_docx(text: str, path: Path) -> None:
    """Минимальный перенос markdown-ответа: заголовки, списки, жирный, курсив."""
    doc = _new_doc()
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if m := re.match(r"^(#{1,6})\s+(.*)", line):
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif m := re.match(r"^\s*[-*•]\s+(.*)", line):
            _add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            _add_inline(doc.add_paragraph(style="List Number"), m.group(1))
        elif re.match(r"^\s*(-{3,}|\*{3,})\s*$", line):
            continue
        else:
            _add_inline(doc.add_paragraph(), line)
    doc.save(str(path))
