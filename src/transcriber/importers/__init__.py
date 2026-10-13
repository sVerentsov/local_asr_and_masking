"""Импорт готовых транскриптов: txt, md, docx, srt, vtt."""

from __future__ import annotations

import io
import re
from pathlib import Path

from transcriber.model import Utterance

SUPPORTED = (".txt", ".md", ".docx", ".srt", ".vtt")

_TS = r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})(?:[.,]\d+)?"
_LINE_TS_SPEAKER = re.compile(rf"^\s*[\[(]?{_TS}[\])]?\s*[-–—]?\s*([^:\[\]]{{1,40}}?)\s*:\s*(.*)$")
_LINE_SPEAKER_TS = re.compile(rf"^\s*([^:\[\]]{{1,40}}?)\s*[\[(]{_TS}[\])]\s*:\s*(.*)$")
_LINE_SPEAKER = re.compile(r"^\s*(?:\*\*)?([^:\[\]*]{1,40}?)(?:\*\*)?\s*:\s*(?:\*\*)?\s*(.+)$")
_LINE_TS_ONLY = re.compile(rf"^\s*[\[(]{_TS}[\])]\s*(.*)$")
_SRT_TIME = re.compile(r"(\d{1,2}):(\d{2}):(\d{2})[.,](\d{1,3})\s*-->\s*(\d{1,2}):(\d{2}):(\d{2})[.,](\d{1,3})")
_VTT_TIME = re.compile(
    r"(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,](\d{1,3})\s*-->\s*(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,](\d{1,3})"
)
_VTT_VOICE = re.compile(r"<v(?:\.[\w.]+)?\s+([^>]+)>")
_TAGS = re.compile(r"<[^>]+>")


class ImportError_(ValueError):
    pass


def decode_text(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1251")
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is None:
            raise ImportError_("Не удалось определить кодировку файла") from None
        return str(best)


def _secs(h, m, s, ms="0") -> float:
    return int(h or 0) * 3600 + int(m) * 60 + int(s) + int((ms or "0").ljust(3, "0")[:3]) / 1000


def _is_speaker(name: str) -> bool:
    name = name.strip()
    if not name or len(name.split()) > 4 or name[0].isdigit() and not name.lower().startswith("спикер"):
        return False
    if re.search(r"[.!?,;\"«»()]", name):
        return False
    return name[0].isupper() or name.lower().startswith(("спикер", "speaker"))


def parse_lines(lines: list[str]) -> list[Utterance]:
    """Реплики «Имя: текст», «[00:12:30] Имя: текст». Без узнаваемой структуры — абзацы без спикеров."""
    lines = [ln.rstrip() for ln in lines]
    nonempty = [ln for ln in lines if ln.strip()]
    if not nonempty:
        return []
    utts: list[Utterance] = []
    structured = 0
    for ln in lines:
        if not ln.strip():
            continue
        u = None
        if m := _LINE_TS_SPEAKER.match(ln):
            if _is_speaker(m.group(4)):
                u = Utterance(m.group(5).strip(), m.group(4).strip(), _secs(m.group(1), m.group(2), m.group(3)))
        if u is None and (m := _LINE_SPEAKER_TS.match(ln)):
            if _is_speaker(m.group(1)):
                u = Utterance(m.group(5).strip(), m.group(1).strip(), _secs(m.group(2), m.group(3), m.group(4)))
        if u is None and (m := _LINE_SPEAKER.match(ln)):
            if _is_speaker(m.group(1)):
                u = Utterance(m.group(2).strip(), m.group(1).strip())
        if u is None and (m := _LINE_TS_ONLY.match(ln)):
            u = Utterance(m.group(4).strip(), None, _secs(m.group(1), m.group(2), m.group(3)))
        if u is not None:
            structured += 1
            utts.append(u)
        elif utts and structured:
            utts[-1].text = (utts[-1].text + " " + ln.strip()).strip()
        else:
            utts.append(Utterance(ln.strip()))
    if structured < max(1, len(nonempty) // 3):
        return _paragraphs(lines)
    return [u for u in utts if u.text]


def _paragraphs(lines: list[str]) -> list[Utterance]:
    out, cur = [], []
    for ln in lines:
        if ln.strip():
            cur.append(ln.strip())
        elif cur:
            out.append(Utterance(" ".join(cur)))
            cur = []
    if cur:
        out.append(Utterance(" ".join(cur)))
    return out


def parse_srt(text: str) -> list[Utterance]:
    out = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()):
        lines = [ln for ln in block.split("\n") if ln.strip()]
        t = next((i for i, ln in enumerate(lines) if _SRT_TIME.search(ln)), None)
        if t is None:
            continue
        m = _SRT_TIME.search(lines[t])
        start, end = _secs(*m.group(1, 2, 3, 4)), _secs(*m.group(5, 6, 7, 8))
        body = _TAGS.sub("", " ".join(lines[t + 1 :])).strip()
        out.append(_speaker_from_body(body, start, end))
    return [u for u in out if u.text]


def parse_vtt(text: str) -> list[Utterance]:
    out = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()):
        lines = [ln for ln in block.split("\n") if ln.strip()]
        t = next((i for i, ln in enumerate(lines) if _VTT_TIME.search(ln)), None)
        if t is None:
            continue
        m = _VTT_TIME.search(lines[t])
        start, end = _secs(*m.group(1, 2, 3, 4)), _secs(*m.group(5, 6, 7, 8))
        raw = " ".join(lines[t + 1 :])
        voice = _VTT_VOICE.search(raw)
        body = _TAGS.sub("", raw).strip()
        if voice:
            out.append(Utterance(body, voice.group(1).strip(), start, end))
        else:
            out.append(_speaker_from_body(body, start, end))
    return [u for u in out if u.text]


def _speaker_from_body(body: str, start: float, end: float) -> Utterance:
    m = _LINE_SPEAKER.match(body)
    if m and _is_speaker(m.group(1)):
        return Utterance(m.group(2).strip(), m.group(1).strip(), start, end)
    return Utterance(body, None, start, end)


def _docx_lines(data: bytes) -> list[str]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        lines.extend(p.text.split("\n"))
        lines.append("")
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(": ".join(c for c in cells if c))
    return lines


def _merge_same_speaker(utts: list[Utterance]) -> list[Utterance]:
    """Субтитры режут речь на куски — склеиваем соседние куски одного спикера."""
    out: list[Utterance] = []
    for u in utts:
        if out and u.speaker is not None and out[-1].speaker == u.speaker:
            out[-1].text += " " + u.text
            out[-1].end = u.end
        else:
            out.append(u)
    return out


def import_bytes(data: bytes, filename: str) -> list[Utterance]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ImportError_(f"Формат {ext or 'без расширения'} не поддерживается")
    if ext == ".docx":
        utts = parse_lines(_docx_lines(data))
    else:
        text = decode_text(data)
        if ext == ".srt":
            utts = _merge_same_speaker(parse_srt(text))
        elif ext == ".vtt":
            utts = _merge_same_speaker(parse_vtt(text))
        else:
            if ext == ".md":
                text = re.sub(r"^#+\s.*$", "", text, flags=re.M)
            utts = parse_lines(text.splitlines())
    if not utts:
        raise ImportError_("В файле не найден текст")
    return utts
